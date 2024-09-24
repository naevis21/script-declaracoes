import os
import pandas as pd
from docx import Document

# Caminhos dos arquivos
csv_path = r'C:\Users\Malu\Desktop\pareceristas\DECLARAÇÕES 2023-2024 (respostas) - Respostas ao formulário 1.csv'
template_path = r'C:\Users\Malu\Desktop\pareceristas\ModeloDeclaracao.docx'
output_folder = r'C:\Users\Malu\Desktop\pareceristas\Declarações'

# Função para substituir placeholders no documento
def replace_placeholders(doc, data):
    placeholders = {
        'NOME COMPLETO (SEM ABREVIATURAS)': data.get('NOME COMPLETO (SEM ABREVIATURAS)', ''),
        'TÍTULO DA OBRA': data.get('TÍTULO DA OBRA', ''),
        'CPF (Ex.: 000.000.000-00)': data.get('CPF (Ex.: 000.000.000-00)', ''),
    }
    
    for paragraph in doc.paragraphs:
        for key, value in placeholders.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, str(value))

# Função para gerar DOCXs a partir do modelo e dados do CSV
def generate_docs_from_csv(csv_path, template_path, output_folder):
    # Verifica se o diretório de saída existe, se não, cria
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    
    # Tenta ler os dados do CSV
    try:
        df = pd.read_csv(csv_path)
    except Exception as e:
        print(f"Erro ao ler o CSV: {e}")
        return

    # Filtra linhas onde a coluna "SITUAÇÃO" está vazia
    filtered_df = df[df['SITUAÇÃO'].isna()]

    for index, row in filtered_df.iterrows():
        data = row.to_dict()
        
        # Carregar o modelo DOCX
        try:
            doc = Document(template_path)
        except Exception as e:
            print(f"Erro ao carregar o modelo DOCX: {e}")
            continue

        # Substituir placeholders
        replace_placeholders(doc, data)

        # Salvar o documento final como DOCX com o nome da coluna "NOME COMPLETO (SEM ABREVIATURAS)"
        name = data['NOME COMPLETO (SEM ABREVIATURAS)'].replace('/', '-')  # Substitui caracteres inválidos
        doc_output_path = os.path.join(output_folder, f"{name}.docx")
        doc.save(doc_output_path)

        print(f"DOCX gerado: {doc_output_path}")

# Gerar DOCXs
generate_docs_from_csv(csv_path, template_path, output_folder)
